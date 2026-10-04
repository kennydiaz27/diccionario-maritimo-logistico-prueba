import openpyxl
import json
from pathlib import Path


# ==========================================
# CONFIGURACIÓN
# ==========================================

CARPETA_EXCEL = Path("excel")
CARPETA_DATA = Path("data")


# Buscar automáticamente el archivo Excel
archivos_excel = list(CARPETA_EXCEL.glob("*.xlsx"))

if not archivos_excel:
    print("❌ No se encontró ningún archivo .xlsx dentro de la carpeta 'excel'.")
    print("Verifica que tu Excel esté dentro de esa carpeta.")
    exit()


archivo_excel = archivos_excel[0]

print(f"📘 Excel encontrado: {archivo_excel.name}")


# Crear carpeta data si no existe
CARPETA_DATA.mkdir(exist_ok=True)

archivo_salida = CARPETA_DATA / "diccionario.json"


# ==========================================
# ABRIR EXCEL
# ==========================================

libro = openpyxl.load_workbook(
    archivo_excel,
    data_only=True
)


print("\n📑 Pestañas encontradas:")

for hoja in libro.sheetnames:
    print(f"   - {hoja}")


# ==========================================
# FUNCIONES AUXILIARES
# ==========================================

def limpiar(valor):
    """
    Convierte cualquier valor vacío en texto vacío.
    """

    if valor is None:
        return ""

    return str(valor).strip()


def convertir_filas(hoja):
    """
    Convierte las filas de una hoja de Excel
    en una lista de diccionarios.
    """

    filas = list(
        hoja.iter_rows(
            values_only=True
        )
    )


    if not filas:
        return []


    # Primera fila = encabezados
    encabezados = [

        limpiar(celda).lower()

        for celda in filas[0]

    ]


    registros = []


    for fila in filas[1:]:

        # Ignorar filas completamente vacías
        if all(
            celda is None
            for celda in fila
        ):
            continue


        registro = {}


        for i, valor in enumerate(fila):

            if i < len(encabezados):

                encabezado = encabezados[i]


                if encabezado:

                    registro[encabezado] = limpiar(valor)


        registros.append(registro)


    return registros


# ==========================================
# CONVERSIÓN DE LAS 3 PESTAÑAS
# ==========================================

datos_finales = []


# ==========================================
# TRANSPORTE MARÍTIMO
# ==========================================

if "Transporte Marítimo" in libro.sheetnames:

    hoja = libro["Transporte Marítimo"]

    registros = convertir_filas(hoja)


    print(
        f"\n🚢 Transporte Marítimo: "
        f"{len(registros)} registros"
    )


    for registro in registros:

        item = {

            "id": registro.get(
                "id",
                ""
            ),

            "nivel": registro.get(
                "nivel",
                registro.get(
                    "nivel de dificultad",
                    ""
                )
            ),

            "termino": registro.get(
                "término (español)",
                registro.get(
                    "termino (español)",
                    ""
                )
            ),

            "ingles": registro.get(
                "versión inglés",
                registro.get(
                    "version inglés",
                    ""
                )
            ),

            "tipo": "Término",

            "categoria": "Transporte Marítimo",

            "subcategoria": registro.get(
                "subcategoría",
                registro.get(
                    "subcategoria",
                    ""
                )
            ),

            "definicion_es": registro.get(
                "definición (español)",
                registro.get(
                    "definicion (español)",
                    ""
                )
            ),

            "definicion_en": registro.get(
                "definición (inglés)",
                registro.get(
                    "definicion (inglés)",
                    ""
                )
            ),

            "fuente": registro.get(
                "fuente",
                ""
            ),

            "vigencia": "",

            "estructura": "estandar"

        }


        datos_finales.append(item)


# ==========================================
# LOGÍSTICA INTERNACIONAL
# ==========================================

if "Logística Internacional" in libro.sheetnames:

    hoja = libro["Logística Internacional"]

    registros = convertir_filas(hoja)


    print(
        f"📦 Logística Internacional: "
        f"{len(registros)} registros"
    )


    for registro in registros:

        item = {

            "id": registro.get(
                "id",
                ""
            ),

            "nivel": registro.get(
                "nivel",
                registro.get(
                    "nivel de dificultad",
                    ""
                )
            ),

            "termino": registro.get(
                "término (español)",
                registro.get(
                    "termino (español)",
                    ""
                )
            ),

            "ingles": registro.get(
                "versión inglés",
                registro.get(
                    "version inglés",
                    ""
                )
            ),

            "tipo": "Término",

            "categoria": "Logística Internacional",

            "subcategoria": registro.get(
                "subcategoría",
                registro.get(
                    "subcategoria",
                    ""
                )
            ),

            "definicion_es": registro.get(
                "definición (español)",
                registro.get(
                    "definicion (español)",
                    ""
                )
            ),

            "definicion_en": registro.get(
                "definición (inglés)",
                registro.get(
                    "definicion (inglés)",
                    ""
                )
            ),

            "fuente": registro.get(
                "fuente",
                ""
            ),

            "vigencia": "",

            "estructura": "estandar"

        }


        datos_finales.append(item)


# ==========================================
# SIGLAS Y ABREVIATURAS
# ==========================================

if "Siglas y Abreviaturas" in libro.sheetnames:

    hoja = libro["Siglas y Abreviaturas"]

    registros = convertir_filas(hoja)


    print(
        f"🔤 Siglas y Abreviaturas: "
        f"{len(registros)} registros"
    )


    for registro in registros:

        # La columna "Categoría" del Excel
        # será utilizada como SUBCATEGORÍA.
        #
        # La categoría principal será siempre:
        # "Siglas y Abreviaturas"

        categoria_sigla = registro.get(
            "categoría",
            registro.get(
                "categoria",
                ""
            )
        )


        item = {

            "id": registro.get(
                "id",
                ""
            ),

            "nivel": registro.get(
                "nivel",
                registro.get(
                    "nivel de dificultad",
                    ""
                )
            ),

            "termino": registro.get(
                "sigla / abreviatura",
                registro.get(
                    "sigla/abreviatura",
                    ""
                )
            ),

            "ingles": registro.get(
                "nombre original (inglés)",
                registro.get(
                    "nombre original (ingles)",
                    ""
                )
            ),

            "tipo": "Sigla / Abreviatura",

            "categoria": "Siglas y Abreviaturas",

            "subcategoria": categoria_sigla,

            "definicion_es": registro.get(
                "traducción / significado (español)",
                registro.get(
                    "traduccion / significado (español)",
                    ""
                )
            ),

            "definicion_en": "",

            "fuente": "",

            "vigencia": "",

            "estructura": "sigla"

        }


        datos_finales.append(item)


# ==========================================
# GUARDAR JSON
# ==========================================

with open(
    archivo_salida,
    "w",
    encoding="utf-8"
) as archivo:

    json.dump(
        datos_finales,
        archivo,
        ensure_ascii=False,
        indent=2
    )


# ==========================================
# RESULTADO
# ==========================================

print("\n==========================================")
print("✅ CONVERSIÓN TERMINADA")
print("==========================================")

print(
    f"📚 Total de registros: "
    f"{len(datos_finales)}"
)

print(
    f"📄 Archivo creado: "
    f"{archivo_salida}"
)

print(
    "\nYa puedes utilizar "
    "'data/diccionario.json' "
    "en la página web."
)